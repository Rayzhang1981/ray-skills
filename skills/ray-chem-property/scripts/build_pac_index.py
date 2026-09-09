# -*- coding: utf-8 -*-
"""构建 DOE PAC/TEEL 本地索引：读取 DOE PAC 数据库 Excel → JSON 索引。

背景（2026-08-15 实测）：
- 官方入口：https://emhub1.energy.gov/pacteel（DOE EM Hub 站点，仅此一处可下载全量数据库）
  ⚠️ 该站从中国大陆网络不可达（直连超时 WinError 10060 / Azure 网关 403），需在美国网络或代理下访问；
- 数据库内容：3000+ 化学品，PAC-1/PAC-2/PAC-3（单位 ppm 或 mg/m3），含 TEEL/AEGL/ERPG 各层级值；
- PAC 层级规则（选值优先级）：Final AEGL(60min) > Interim AEGL(60min) > ERPG > TEEL；
- 替代源（国内可达）：NOAA CAMEO Chemicals 数据自带 PAC 值（skill 已有 CAMEO fetch 链路可覆盖部分）。

用法（拿到 Excel 后）：
  python scripts/build_pac_index.py --xlsx PAC_Rev.30.xlsx            # 解析并构建 → ../data/PAC索引.json
  python scripts/build_pac_index.py --xlsx PAC_Rev.30.xlsx --query 67-64-1   # 只查一个 CAS

输出：data/PAC索引.json
  {"67-64-1": {"name": "Acetone", "unit": "ppm", "pac1": "…", "pac2": "…", "pac3": "…",
               "teel0": "…", "teel1": "…", "teel2": "…", "teel3": "…", "note": ""}, ...}
字段键与 field_schema.json 的 pac1/2/3、teel1/2/3 对齐。
"""
import argparse
import json
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.normpath(os.path.join(HERE, "..", "data"))

# 列名模糊匹配（DOE Excel 各版本列名略有差异，按关键词识别）
COL_HINTS = {
    "cas": ["cas", "cas number", "cas no", "casrn"],
    "name": ["chemical", "chemical name", "name", "compound"],
    "pac1": ["pac-1", "pac1", "pac 1", "pacs"],
    "pac2": ["pac-2", "pac2", "pac 2"],
    "pac3": ["pac-3", "pac3", "pac 3"],
    "teel0": ["teel-0", "teel0", "teel 0"],
    "teel1": ["teel-1", "teel1", "teel 1"],
    "teel2": ["teel-2", "teel2", "teel 2"],
    "teel3": ["teel-3", "teel3", "teel 3"],
    "unit_col": ["units", "unit"],
    "source": ["source of pacs", "source", "basis"],
}


def find_col(header_row, hint_keys):
    """在表头行找匹配列，返回列索引或 None。"""
    for i, cell in enumerate(header_row):
        if cell is None:
            continue
        c = str(cell).strip().lower()
        for h in hint_keys:
            if c == h or c.startswith(h):
                return i
    return None


def norm(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.upper() in ("NR", "NA", "N/A", "-", "--", "U", "ND"):
        return None
    return s


def parse_xlsx(xlsx_path):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.worksheets[0]  # 数据通常在第一个 sheet
    rows = ws.iter_rows(values_only=True)

    # 多级表头合并：找含 CAS 关键词的行作 R1，取下一行作 R2（单位/子标题），逐列拼接
    header_r1 = None
    for _ in range(10):
        try:
            row = next(rows)
        except StopIteration:
            break
        if any(row and ("cas" in str(c).strip().lower() for c in row if c)):
            header_r1 = row
            break
    if header_r1 is None:
        raise SystemExit("[pac] 前 10 行未找到含 CAS 的表头行——请检查 Excel 结构（可能需手动指定 sheet）")
    try:
        header_r2 = next(rows)
    except StopIteration:
        header_r2 = ()
    header = []
    for i in range(max(len(header_r1), len(header_r2))):
        a = str(header_r1[i]) if i < len(header_r1) and header_r1[i] is not None else ""
        b = str(header_r2[i]) if i < len(header_r2) and header_r2[i] is not None else ""
        header.append(f"{a} {b}".strip().lower())

    idx = {k: find_col(header, v) for k, v in COL_HINTS.items()}
    missing = [k for k, v in idx.items() if v is None and k not in ("teel0", "teel1", "teel2", "teel3")]
    if missing:
        print(f"[pac] 警告：以下列未识别（跳过）: {missing}")
        print(f"[pac] 合并表头: {[h[:30] for h in header if h][:22]}")

    recs = {}
    for row in rows:
        cas = norm(row[idx["cas"]]) if idx.get("cas") is not None else None
        if not cas or not re.match(r"^\d{2,7}-\d{2}-\d$", cas):
            continue
        name = norm(row[idx["name"]]) if idx.get("name") is not None else ""
        rec = {"name": name or "", "unit": ""}
        # 单位列（col19 Units）优先；否则按 PAC 值含字母判断
        if idx.get("unit_col") is not None:
            u = norm(row[idx["unit_col"]])
            if u:
                rec["unit"] = u.lower()
        for k in ("pac1", "pac2", "pac3", "teel0", "teel1", "teel2", "teel3"):
            v = norm(row[idx[k]]) if idx.get(k) is not None else None
            if v and k.startswith("pac") and rec["unit"] == "":
                rec["unit"] = "mg/m3" if re.search(r"[a-zA-Z]", v) else "ppm"
            rec[k] = v
        # 来源列（Source of PACs）保留为备注
        if idx.get("source") is not None:
            src = norm(row[idx["source"]])
            if src:
                rec["note"] = src[:60]
        recs[cas] = rec

    wb.close()
    return recs


def main():
    ap = argparse.ArgumentParser(description="构建 DOE PAC/TEEL 本地索引")
    ap.add_argument("--xlsx", required=True, help="DOE PAC 数据库 Excel（从 emhub1.energy.gov/pacteel 下载）")
    ap.add_argument("--query", help="只查询某个 CAS 并打印")
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "PAC索引.json"))
    args = ap.parse_args()

    if not os.path.exists(args.xlsx):
        raise SystemExit(f"[pac] 找不到 {args.xlsx}\n  下载提示：PAC/TEEL 数据库需在能访问 emhub1.energy.gov 的网络下载（国内不可达），下载后传入 --xlsx")

    recs = parse_xlsx(args.xlsx)
    print(f"[pac] 解析到 {len(recs)} 个化学品")

    if args.query:
        q = args.query.strip()
        hit = recs.get(q)
        if not hit:
            for cas, r in recs.items():
                if q.lower() in (r["name"] or "").lower():
                    hit, q = r, cas
                    break
        if hit:
            print(f"  {q} {hit['name']} ({hit['unit']})")
            for k in ("pac1", "pac2", "pac3", "teel0", "teel1", "teel2", "teel3"):
                print(f"    {k}: {hit[k]}")
        else:
            print(f"  CAS {q} 未收录")
        return

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(recs, f, ensure_ascii=False, indent=1)
    print(f"[pac] 已保存 {args.out}")
    n = sum(1 for r in recs.values() if r.get("pac3"))
    print(f"[pac] PAC-3 有值: {n}")


if __name__ == "__main__":
    main()
