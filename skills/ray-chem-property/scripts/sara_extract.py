import openpyxl, math, json, os
SRC = os.environ['SARA_SRC']
OUT = os.environ['SARA_OUT']
LOG = os.environ.get('SARA_LOG', './output/sara_log.txt')
def log(m):
    with open(LOG, 'a', encoding='utf-8') as f:
        f.write(m + '\n')
log('LOAD START')
wb = openpyxl.load_workbook(SRC, data_only=True)
log('LOADED')
ws = wb['physical-properties']
groups = {'hvap':7,'vp':16,'cpg':25,'cpl':34,'denl':43,'visg':52,'visl':61}
def vp(c,T):
    try: return math.exp(c[0]+c[1]/T+c[2]*math.log(T)+c[3]*T**c[4])
    except: return None
idx = {}
for r in range(6, ws.max_row+1):
    n = ws.cell(row=r,column=1).value
    if not n or not isinstance(n,str): continue
    rec={'name':n.strip(),'formula':str(ws.cell(row=r,column=2).value or '').strip() or None,
         'tc_c':ws.cell(row=r,column=3).value,'pc_bar':ws.cell(row=r,column=4).value,
         'mw':ws.cell(row=r,column=5).value,'nbp_c':ws.cell(row=r,column=6).value}
    for g,c0 in groups.items():
        code=ws.cell(row=r,column=c0).value
        if isinstance(code,(int,float)):
            rec[g+'_coef']={'code':int(code),'tmin':ws.cell(row=r,column=c0+1).value,
                            'tmax':ws.cell(row=r,column=c0+2).value,
                            'c':[ws.cell(row=r,column=c).value for c in range(c0+3,c0+9)]}
    if 'vp_coef' in rec:
        vc=rec['vp_coef']['c']
        for T in (20,25,60):
            pa=vp(vc,273.15+T)
            if pa is not None: rec['vp_%dc_pa'%T]=round(pa,1)
    idx[rec['name'].lower()]=rec
log('LOOP DONE %d' % len(idx))
wb.close()
os.makedirs(os.path.dirname(OUT),exist_ok=True)
json.dump(idx, open(OUT,'w',encoding='utf-8'), ensure_ascii=False, indent=1)
log('DUMP DONE')
