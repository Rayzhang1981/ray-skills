# -*- coding: utf-8 -*-
"""数据源适配器：天天基金持仓 Excel → 标准化持仓 JSON。

本文件是「数据源适配层」的一个 adapter。职责单一：把天天基金 Excel
映射成标准化持仓 JSON（见 SKILL.md「标准化持仓 JSON」章节），下游分析零依赖。

用法：
    python excel.py <excel路径> <日期YYYY-MM-DD> [输出json路径]

示例：
    python excel.py <path/to/export.xlsx> <快照日期如2024-01-01>

输出 _portfolio_data.json，字段：
    mm_funds: [{name, yield_rate, shares, unpaid, cumulative_gain}]
    funds: [{name, type_info, acls, amount, status, gain, gpct, gpct_s, nav}]
    mm_total / fund_total / grand_total / class_totals / class_gains / date
"""
import openpyxl
import json
import os
import sys

# ===== 参数 =====
if len(sys.argv) < 3:
    print(__doc__)
    sys.exit(1)

excel_path = sys.argv[1]
date_str = sys.argv[2]
# 默认输出到技能包外的数据目录（家庭财务敏感数据不进技能包——发布安全 + 换环境可移植）；第 3 参数可覆盖
_DEFAULT_DATA_DIR = os.path.join(os.path.expanduser('~'), '.workbuddy', 'data', 'ray-household-finance')
os.makedirs(_DEFAULT_DATA_DIR, exist_ok=True)
out_path = sys.argv[3] if len(sys.argv) > 3 else os.path.join(_DEFAULT_DATA_DIR, '_portfolio_data.json')

wb = openpyxl.load_workbook(excel_path, data_only=True)

# ===== 活期宝 sheet =====
# sheet 名可能为「活期宝 / Sheet1 / Sheet2 ...」——按位置取：第一个 sheet 为活期宝
# （天天基金导出：活期宝在前、基金在后；0823 版导出名为 Sheet1/Sheet2）
_ws0 = wb.worksheets[0]
if '简称' in str(_ws0.cell(1, 1).value or ''):
    ws_mm = _ws0
else:
    ws_mm = wb.worksheets[1]
    # 找不到活期宝表头时，回退按名字匹配
    for _n in ['活期宝', '货币', '余额宝']:
        if _n in wb.sheetnames:
            ws_mm = wb[_n]
            break
mm_funds = []
mm_total = 0.0
for row in ws_mm.iter_rows(min_row=1, max_row=ws_mm.max_row, values_only=True):
    non_none = [v for v in row if v is not None]
    if len(non_none) < 5:
        continue
    name = row[0]
    if not isinstance(name, str):
        continue
    if '简称' in name or '操作' in name:
        continue
    # 份额：'可用 / 可取 / 总份额'，取最后一段
    parts = str(row[2]).replace(',', '').split('/') if row[2] else ['0']
    total_shares = float(parts[-1].strip())
    unpaid = float(row[3]) if row[3] else 0.0
    cumulative = float(row[4]) if row[4] else 0.0
    mm_funds.append({
        'name': name,
        'yield_rate': str(row[1]) if row[1] else '',
        'shares': total_shares,
        'unpaid': unpaid,
        'cumulative_gain': cumulative
    })
    mm_total += total_shares

# ===== 基金 sheet =====
# 每只基金占一个块：名称行(含6位代码) + 类型行 + 金额行 + 状态行 + 盈亏行 + 收益率行
# 状态值「有在途交易/到期/未到期/将到期」使块高为 7；否则为 6
# sheet 名可能为「基金 / Sheet2 ...」——取另一个 sheet（活期宝之外的）
ws_fund = None
for _n in ['基金', 'fund', 'Fund']:
    if _n in wb.sheetnames:
        ws_fund = wb[_n]
        break
if ws_fund is None:
    ws_fund = wb.worksheets[1] if len(wb.worksheets) > 1 else wb.worksheets[0]
all_rows = [row[0] for row in ws_fund.iter_rows(min_row=1, max_row=ws_fund.max_row, values_only=True)]

STATUS_VALUES = ['有在途交易', '到期', '未到期', '将到期']

funds = []
i = 4
while i < len(all_rows):
    name = all_rows[i]
    if name is None or not isinstance(name, str):
        i += 1
        continue
    # 跳过表头/分隔行
    if any(k in name for k in ['买入卖出', '按产品', '金额', '持仓收益', '操作']):
        i += 1
        continue

    ft = str(all_rows[i + 1]) if all_rows[i + 1] else ''
    amt = float(all_rows[i + 2]) if all_rows[i + 2] else 0.0

    nv = all_rows[i + 3]
    status = ''
    gain = 0.0
    gpct_s = ''

    if isinstance(nv, str) and nv in STATUS_VALUES:
        # 有状态值：块高 7
        status = nv
        gain = float(all_rows[i + 4]) if all_rows[i + 4] and all_rows[i + 4] != '--' else 0.0
        gv = all_rows[i + 5]
        gpct_s = str(gv) if gv is not None and gv != '--' else '--'
        i += 7
    else:
        # 无状态值：块高 6
        gain = float(nv) if nv and nv != '--' else 0.0
        gv = all_rows[i + 4]
        gpct_s = str(gv) if gv is not None and gv != '--' else '--'
        i += 6

    # 资产分类
    if '债券型' in ft:
        acls = 'bond'
    elif '指数型' in ft:
        acls = 'index'
    elif 'QDII' in ft:
        acls = 'qdii'
    elif '混合型' in ft:
        acls = 'mixed'
    else:
        acls = 'other'

    # 提取净值
    nav = ''
    if '最新净值' in ft:
        nav = ft.split('：')[-1].split('（')[0].split('(')[0].strip()

    try:
        gpct = float(gpct_s) if gpct_s not in ('', '--') else 0.0
    except ValueError:
        gpct = 0.0

    funds.append({
        'name': name,
        'type_info': ft,
        'acls': acls,
        'amount': round(amt, 2),
        'status': status,
        'gain': round(gain, 2),
        'gpct': round(gpct, 6),
        'gpct_s': gpct_s,
        'nav': nav
    })

funds.sort(key=lambda x: x['amount'], reverse=True)
fund_total = sum(f['amount'] for f in funds)
grand_total = mm_total + fund_total

# 类别汇总
ct = {}
cg = {}
for f in funds:
    ct[f['acls']] = ct.get(f['acls'], 0) + f['amount']
    cg[f['acls']] = cg.get(f['acls'], 0) + f['gain']
ct['mm'] = mm_total
cg['mm'] = sum(mf['cumulative_gain'] for mf in mm_funds)

data = {
    'mm_funds': mm_funds,
    'funds': funds,
    'mm_total': round(mm_total, 2),
    'fund_total': round(fund_total, 2),
    'grand_total': round(grand_total, 2),
    'class_totals': {k: round(v, 2) for k, v in ct.items()},
    'class_gains': {k: round(v, 2) for k, v in cg.items()},
    'date': date_str
}

with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print(f'OK: {len(funds)} funds, {len(mm_funds)} MM, total={grand_total:,.2f}')
print(f'Classes: { {k: round(v, 0) for k, v in ct.items()} }')
