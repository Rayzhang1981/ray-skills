# -*- coding: utf-8 -*-
"""演讲稿 MD 同源导出（v3.20.0 新增）

为什么要它：SKILL.md 要求"生成 PPTX 时同步导出演讲稿 MD"，但"另写一遍"= 两处真相，
改一处必漂移。正解是从 PPTX 的 notes_slide 回读 —— 单源，改口播稿就改 PPTX 备注层再重跑。

用法：
    python scripts/export_speech.py <文件.pptx> <输出.md> --titles "封面,结论先行,措施一,..."
    python scripts/export_speech.py <文件.pptx> <输出.md> --titles-file titles.txt

    --titles  逗号分隔的页标题，顺序 = 页序；缺省则页标题写"第 N 页"
    页标题不从 deck 里猜 —— 标题散在 add_title_bar 的文本框里，猜错比留空更坏。
"""
import sys, argparse


def main():
    ap = argparse.ArgumentParser(description="从 PPTX 演讲者备注导出逐页演讲稿 MD")
    ap.add_argument("pptx")
    ap.add_argument("out")
    ap.add_argument("--titles", default="", help="逗号分隔的页标题")
    ap.add_argument("--titles-file", default="", help="每行一个页标题的文件")
    ap.add_argument("--intro", default="", help="写进文首的一句话说明")
    a = ap.parse_args()

    from pptx import Presentation
    titles = []
    if a.titles_file:
        titles = [l.strip() for l in open(a.titles_file, encoding="utf-8") if l.strip()]
    elif a.titles:
        titles = [t.strip() for t in a.titles.split(",")]
        if len(titles) == 1:
            titles = [t.strip() for t in a.titles.split("，")]

    prs = Presentation(a.pptx)
    lines = ["# 演讲稿", ""]
    if a.intro:
        lines += [f"> {a.intro}", ""]
    lines += [f"> 与 `{a.pptx}` 逐页一一对应（共 {len(prs.slides._sldIdLst)} 页）。",
              "> 本稿自 PPTX 演讲者备注导出；修改请回 PPTX 备注层后重新导出，避免两处漂移。", "", "---", ""]

    n = 0
    for i, slide in enumerate(prs.slides):
        n += 1
        t = ""
        if slide.has_notes_slide:
            t = slide.notes_slide.notes_text_frame.text.strip()
        title = titles[i] if i < len(titles) else f"第 {n} 页"
        lines += [f"## 第 {n} 页：{title}", "", t if t else "（本页无口播稿）", "", "---", ""]

    with open(a.out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"[OK] {n} 页演讲稿 -> {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
