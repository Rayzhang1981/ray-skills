#!/usr/bin/env python3
"""
ray-translate 精细残留诊断 — 通用版
使用白名单过滤出真正需要翻译的残留英文
用法：
    python extract_real_remaining.py <翻译后文件.xlsm> [--sheets Sheet1 Sheet2 ...]
"""
import openpyxl
import re
import sys
from collections import Counter

# ============================================================
# 白名单 — 以下内容不应翻译
# 按需增删，尤其注意项目特有的缩写/变量名
# ============================================================
KEEP_ENGLISH = {
    # ---- 单位符号 ----
    'kPa', 'ppm', 'ppmv', 'bar', 'min', 'sec', 'BTU', 'Btu', 'kcal',
    'atm', 'psi', 'mmHg', 'vol', 'gal', 'liter', 'mol', 'g', 'cal',
    'kg', 'm3', 'kW', 'kWatt', 'kgsec', 'lbmin', 'kcalmin',
    'kgsq', 'calg', 'volp', 'msec',
    # ---- 专业缩写 ----
    'NFPA', 'ERPG', 'ERPG1', 'ERPG2', 'ERPG3', 'LFL', 'UFL',
    'PFD', 'IEF', 'CAS', 'CCPS', 'CHEF', 'MTTR', 'MTBF',
    'BLEVE', 'VCE', 'TNT', 'LOPA', 'BPCS', 'SIF', 'IPL', 'IPLs',
    'PFDavg', 'PFDs', 'DMSO', 'FMEA', 'HAZOP',
    # ---- 公式变量名 ----
    'Psat', 'DHV', 'DHr', 'VPES', 'VEquip', 'QPV', 'QChem',
    'DHR', 'TNR', 'Tmax', 'TMax', 'Pmax', 'TBurst',
    'Ppad', 'TRef', 'Tcoolant', 'tMR', 'rRef',
    'XLFL', 'XEE', 'XLC', 'CLFL', 'kgTNTeq',
    'PES', 'TAmb', 'T0', 'P0', 'PA', 'PB', 'FV', 'Fv',
    'Mw', 'DE', 'CS', 'TB', 'Cd',
    # ---- 机构/人名 ----
    'AIChE', 'Dow', 'Baker', 'Strehlow', 'Tang', 'Britter', 'McQuaid',
    'Crowl', 'Louvar', 'Mannan', 'Perry', 'RAST', 'Lee',
    'Center', 'Society',
    # ---- 导航/UI ----
    'Go', 'Top',
}


def scan_real_remaining(filepath, target_sheets=None, keep_english=None):
    if keep_english is None:
        keep_english = KEEP_ENGLISH

    wb = openpyxl.load_workbook(filepath, keep_vba=True, data_only=False)
    sheets = target_sheets if target_sheets else wb.sheetnames

    total_cells = 0

    for sname in sheets:
        ws = wb[sname]
        cells_with_real_eng = []

        for row_idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column), 1
        ):
            for col_idx, cell in enumerate(row, 1):
                if cell.value is None:
                    continue
                val = cell.value
                if isinstance(val, str) and not val.startswith("="):
                    # 提取 >=2 字母的英文
                    eng_words = re.findall(r'\b[A-Za-z]{2,}(?:/[A-Za-z]{2,})*\b', val)
                    meaningful = [
                        w for w in eng_words
                        if w not in keep_english and len(w) >= 2
                    ]
                    if meaningful:
                        cells_with_real_eng.append({
                            'row': row_idx,
                            'col': col_idx,
                            'full_value': val,
                            'mean_eng': meaningful,
                        })

        if cells_with_real_eng:
            print(f"\n{'=' * 80}")
            print(f"Sheet: {sname} -- {len(cells_with_real_eng)} cells with REAL English")
            print(f"{'=' * 80}")
            for c in cells_with_real_eng[:20]:
                print(f"\n  Cell [{c['row']},{c['col']}]:")
                print(f"    ENG: {c['mean_eng'][:15]}")
                print(f"    VAL: {c['full_value'][:250]}")
            if len(cells_with_real_eng) > 20:
                print(f"  ... and {len(cells_with_real_eng) - 20} more cells")
            total_cells += len(cells_with_real_eng)

    print(f"\n{'=' * 80}")
    print(f"TOTAL cells needing translation: {total_cells}")
    return total_cells


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python extract_real_remaining.py <翻译后文件.xlsm> [--sheets Sheet1 Sheet2 ...]")
        sys.exit(1)

    filepath = sys.argv[1]
    target_sheets = None

    if "--sheets" in sys.argv:
        idx = sys.argv.index("--sheets")
        target_sheets = sys.argv[idx + 1:]

    scan_real_remaining(filepath, target_sheets)